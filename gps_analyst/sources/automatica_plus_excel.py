from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gps_analyst.models.gps import (
    GpsDetailEvent,
    GpsWorkbookData,
    VehicleDay,
    VehicleTotal,
)


_REQUIRED_SHEETS = {"Totales", "SubTotales", "Detalle"}

_EXPORT_RANGE_RE = re.compile(
    r"INICIO:\s*(\d{2}-\d{2}-\d{4}\s+\d{2}:\d{2}:\d{2})"
    r"\s+FIN:\s*(\d{2}-\d{2}-\d{4}\s+\d{2}:\d{2}:\d{2})",
    re.IGNORECASE,
)

_BLOCK_RE = re.compile(
    r"^Inicio de bloque:\s*(\d{2}-\d{2}-\d{4})$",
    re.IGNORECASE,
)


class AutomaticaPlusExcelSource:
    """
    Lector de exportaciones XLSX de Automatica PLUS.

    La fuente se limita a transformar el formato del proveedor a un modelo
    GPS común. No contiene reglas laborales ni lógica de otros sistemas.
    """

    def load(self, path: str | Path) -> GpsWorkbookData:
        source_path = Path(path)

        if not source_path.is_file():
            raise FileNotFoundError(source_path)

        if source_path.suffix.lower() != ".xlsx":
            raise ValueError(
                f"Formato no soportado: {source_path.suffix or '(sin extensión)'}"
            )

        workbook = load_workbook(
            source_path,
            data_only=True,
            read_only=False,
        )

        try:
            missing = _REQUIRED_SHEETS.difference(workbook.sheetnames)
            if missing:
                raise ValueError(
                    "Excel Automatica PLUS incompleto. "
                    f"Faltan hojas: {', '.join(sorted(missing))}"
                )

            export_start, export_end = self._parse_export_range(
                workbook["Totales"]["A2"].value
            )

            totals = self._parse_totals(workbook["Totales"])
            summaries = self._parse_subtotals(workbook["SubTotales"])
            event_map = self._parse_detail(workbook["Detalle"])

            days: list[VehicleDay] = []

            for key, summary in summaries.items():
                events = self._normalize_events_for_summary(
                    summary=summary,
                    events=tuple(event_map.get(key, ())),
                )

                days.append(
                    VehicleDay(
                        vehicle=summary["vehicle"],
                        block_date=summary["block_date"],
                        start_at=summary["start_at"],
                        end_at=summary["end_at"],
                        jornada_seconds=summary["jornada_seconds"],
                        driving_seconds=summary["driving_seconds"],
                        stop_seconds=summary["stop_seconds"],
                        max_speed_kmh=summary["max_speed_kmh"],
                        distance_km=summary["distance_km"],
                        events=events,
                    )
                )

            days.sort(key=lambda item: (item.block_date, item.vehicle))

            issues = self._validate(
                totals=totals,
                days=days,
            )

            return GpsWorkbookData(
                source_path=source_path,
                export_start=export_start,
                export_end=export_end,
                totals=tuple(totals),
                days=tuple(days),
                validation_issues=tuple(issues),
            )
        finally:
            workbook.close()

    def _parse_export_range(
        self,
        value: Any,
    ) -> tuple[datetime | None, datetime | None]:
        text = self._text(value)

        if not text:
            return None, None

        match = _EXPORT_RANGE_RE.search(text)
        if not match:
            return None, None

        return (
            self._parse_datetime(match.group(1)),
            self._parse_datetime(match.group(2)),
        )

    def _parse_totals(self, sheet) -> list[VehicleTotal]:
        totals: list[VehicleTotal] = []

        for row in range(5, sheet.max_row + 1):
            vehicle = self._text(sheet.cell(row, 1).value)

            if not vehicle:
                continue

            if vehicle.lower().startswith("inicio de bloque:"):
                continue

            totals.append(
                VehicleTotal(
                    vehicle=vehicle,
                    jornada_seconds=self._duration_seconds(
                        sheet.cell(row, 2).value
                    )
                    or 0,
                    driving_seconds=self._duration_seconds(
                        sheet.cell(row, 3).value
                    )
                    or 0,
                    stop_seconds=self._duration_seconds(
                        sheet.cell(row, 4).value
                    )
                    or 0,
                    max_speed_kmh=self._number(
                        sheet.cell(row, 5).value
                    )
                    or 0.0,
                    distance_km=self._number(
                        sheet.cell(row, 6).value
                    )
                    or 0.0,
                )
            )

        return totals

    def _parse_subtotals(self, sheet) -> dict[tuple[date, str], dict[str, Any]]:
        summaries: dict[tuple[date, str], dict[str, Any]] = {}
        current_date: date | None = None

        for row in range(5, sheet.max_row + 1):
            first = self._text(sheet.cell(row, 1).value)

            if not first:
                continue

            block_date = self._block_date(first)
            if block_date is not None:
                current_date = block_date
                continue

            if current_date is None:
                raise ValueError(
                    f"Vehículo encontrado sin fecha de bloque en SubTotales, fila {row}"
                )

            start_at = self._optional_datetime(sheet.cell(row, 2).value)
            end_at = self._optional_datetime(sheet.cell(row, 3).value)

            key = (current_date, first)

            if key in summaries:
                raise ValueError(
                    "Resumen duplicado para "
                    f"{current_date.isoformat()} / {first}"
                )

            summaries[key] = {
                "vehicle": first,
                "block_date": current_date,
                "start_at": start_at,
                "end_at": end_at,
                "jornada_seconds": self._duration_seconds(
                    sheet.cell(row, 4).value
                )
                or 0,
                "driving_seconds": self._duration_seconds(
                    sheet.cell(row, 5).value
                )
                or 0,
                "stop_seconds": self._duration_seconds(
                    sheet.cell(row, 6).value
                )
                or 0,
                "max_speed_kmh": self._number(
                    sheet.cell(row, 7).value
                )
                or 0.0,
                "distance_km": self._number(
                    sheet.cell(row, 8).value
                )
                or 0.0,
            }

        return summaries

    def _parse_detail(
        self,
        sheet,
    ) -> dict[tuple[date, str], list[GpsDetailEvent]]:
        events: dict[tuple[date, str], list[GpsDetailEvent]] = defaultdict(list)

        current_date: date | None = None
        current_vehicle: str | None = None

        for row in range(5, sheet.max_row + 1):
            first = self._text(sheet.cell(row, 1).value)
            stop_value = sheet.cell(row, 2).value
            restart_value = sheet.cell(row, 3).value

            if first:
                block_date = self._block_date(first)

                if block_date is not None:
                    current_date = block_date
                    current_vehicle = None
                    continue

                if current_date is None:
                    raise ValueError(
                        f"Vehículo encontrado sin fecha de bloque en Detalle, fila {row}"
                    )

                current_vehicle = first
                events.setdefault((current_date, current_vehicle), [])
                continue

            stop_text = self._text(stop_value)
            restart_text = self._text(restart_value)

            if not stop_text and not restart_text:
                continue

            if current_date is None or current_vehicle is None:
                raise ValueError(
                    f"Fila de detalle sin vehículo activo, fila {row}"
                )

            if stop_text == "??" and restart_text not in {"", "??"}:
                kind = "opening"
            elif stop_text not in {"", "??"} and restart_text == "??":
                kind = "closing"
            elif stop_text not in {"", "??"} and restart_text not in {"", "??"}:
                kind = "segment"
            else:
                raise ValueError(
                    f"Combinación PARA/ARRANCA no reconocida en fila {row}: "
                    f"{stop_text!r} / {restart_text!r}"
                )

            raw_speed = self._number(sheet.cell(row, 7).value)
            raw_distance = self._number(sheet.cell(row, 8).value)

            if kind == "opening":
                normalized_speed = 0.0
                normalized_distance = 0.0
                driving_seconds = None
                stop_seconds = None
            else:
                normalized_speed = raw_speed
                normalized_distance = raw_distance
                driving_seconds = self._duration_seconds(
                    sheet.cell(row, 5).value
                )
                stop_seconds = (
                    self._duration_seconds(sheet.cell(row, 6).value)
                    if kind == "segment"
                    else None
                )

            address_value = self._text(sheet.cell(row, 9).value)
            stop_type, address = self._split_stop_address(address_value)

            hyperlink = sheet.cell(row, 10).hyperlink
            map_target = hyperlink.target if hyperlink is not None else None

            event = GpsDetailEvent(
                vehicle=current_vehicle,
                block_date=current_date,
                kind=kind,
                stop_at=(
                    None
                    if stop_text in {"", "??"}
                    else self._parse_clock_time(stop_value)
                ),
                restart_at=(
                    None
                    if restart_text in {"", "??"}
                    else self._parse_clock_time(restart_value)
                ),
                driving_seconds=driving_seconds,
                stop_seconds=stop_seconds,
                max_speed_kmh=normalized_speed,
                distance_km=normalized_distance,
                raw_max_speed_kmh=raw_speed,
                raw_distance_km=raw_distance,
                stop_type=stop_type,
                address=address,
                map_target=map_target,
            )

            events[(current_date, current_vehicle)].append(event)

        return events

    @staticmethod
    def _distance_rounding_tolerance(parts: int) -> float:
        return (max(parts, 0) + 1) * 0.05 + 0.001

    @staticmethod
    def _clock_at_or_after(
        reference: datetime,
        value: time,
    ) -> datetime:
        candidate = datetime.combine(reference.date(), value)

        if candidate < reference:
            candidate += timedelta(days=1)

        return candidate

    def _normalize_events_for_summary(
        self,
        summary: dict[str, Any],
        events: tuple[GpsDetailEvent, ...],
    ) -> tuple[GpsDetailEvent, ...]:
        if not events:
            return events

        start_at = summary["start_at"]
        end_at = summary["end_at"]

        if start_at is None or end_at is None:
            return events

        # Some Automatica PLUS exports omit the explicit opening row and instead
        # emit a zero-duration segment at the exact start of the day. Treat it
        # as the opening marker while preserving the raw distance/speed fields.
        first = events[0]
        if (
            first.kind == "segment"
            and first.stop_at is not None
            and first.restart_at is not None
            and first.stop_at == start_at.time()
            and first.restart_at == first.stop_at
            and (first.driving_seconds or 0) == 0
            and (first.stop_seconds or 0) == 0
        ):
            first = replace(
                first,
                kind="opening",
                stop_at=None,
                driving_seconds=None,
                stop_seconds=None,
                max_speed_kmh=0.0,
                distance_km=0.0,
            )
            events = (first, *events[1:])

        # Detail duration fields can occasionally be carried from a neighbouring
        # row. Rebuild them from the explicit time boundaries, but only accept
        # the rebuilt values when the complete reconstructed totals match the
        # SubTotales summary exactly.
        cursor = start_at
        rebuilt: list[GpsDetailEvent] = []
        temporal_ok = True

        for event in events:
            if event.kind == "opening":
                rebuilt.append(event)
                continue

            if event.stop_at is None:
                temporal_ok = False
                break

            stop_at = self._clock_at_or_after(cursor, event.stop_at)
            driving_seconds = int((stop_at - cursor).total_seconds())

            if event.kind == "segment":
                if event.restart_at is None:
                    temporal_ok = False
                    break

                restart_at = self._clock_at_or_after(
                    stop_at,
                    event.restart_at,
                )
                stop_seconds = int(
                    (restart_at - stop_at).total_seconds()
                )
                cursor = restart_at
            else:
                stop_seconds = None
                cursor = stop_at

            rebuilt.append(
                replace(
                    event,
                    driving_seconds=driving_seconds,
                    stop_seconds=stop_seconds,
                )
            )

        if temporal_ok and len(rebuilt) == len(events):
            rebuilt_driving = sum(
                event.driving_seconds or 0
                for event in rebuilt
                if event.kind != "opening"
            )
            rebuilt_stop = sum(
                event.stop_seconds or 0
                for event in rebuilt
                if event.kind == "segment"
            )

            if (
                rebuilt_driving == summary["driving_seconds"]
                and rebuilt_stop == summary["stop_seconds"]
            ):
                events = tuple(rebuilt)

        # Explicit opening rows already have their carried distance normalized
        # by _parse_detail. Distance reconciliation below is only for exports
        # where the opening row itself is omitted.
        if events[0].kind == "opening":
            return events

        real_events = [
            event
            for event in events
            if event.kind != "opening"
        ]

        if not real_events:
            return events

        summary_distance = float(summary["distance_km"])
        computed_distance = sum(
            event.distance_km or 0.0
            for event in real_events
        )
        tolerance = self._distance_rounding_tolerance(
            len(real_events)
        )

        if computed_distance - summary_distance <= tolerance:
            return events

        first = real_events[0]

        if (
            first.kind not in {"segment", "closing"}
            or first.distance_km is None
            or first.stop_at is None
        ):
            return events

        remaining_distance = sum(
            event.distance_km or 0.0
            for event in real_events[1:]
        )
        inferred_distance = round(
            summary_distance - remaining_distance,
            3,
        )

        if inferred_distance < -tolerance:
            return events

        inferred_distance = max(0.0, inferred_distance)

        if inferred_distance > first.distance_km + tolerance:
            return events

        # Guard the correction with the temporal boundary of the first trip.
        first_stop_at = self._clock_at_or_after(
            start_at,
            first.stop_at,
        )
        expected_first_driving = int(
            (first_stop_at - start_at).total_seconds()
        )

        if (
            first.driving_seconds is not None
            and first.driving_seconds != expected_first_driving
        ):
            return events

        summary_max_speed = float(summary["max_speed_kmh"])

        if expected_first_driving > 0 and summary_max_speed > 0:
            inferred_average = (
                inferred_distance
                * 3600
                / expected_first_driving
            )

            if inferred_average > summary_max_speed + 0.5:
                return events

        normalized_speed = first.max_speed_kmh
        if (
            normalized_speed is not None
            and summary_max_speed > 0
            and normalized_speed > summary_max_speed + 0.001
        ):
            normalized_speed = None

        normalized_first = replace(
            first,
            distance_km=inferred_distance,
            max_speed_kmh=normalized_speed,
        )

        return (normalized_first, *events[1:])

    def _validate(
        self,
        totals: list[VehicleTotal],
        days: list[VehicleDay],
    ) -> list[str]:
        issues: list[str] = []

        for day in days:
            label = f"{day.block_date.isoformat()} / {day.vehicle}"

            if day.jornada_seconds != day.driving_seconds + day.stop_seconds:
                issues.append(
                    f"{label}: JORNADA != H.MOTOR + H.STOP"
                )

            if not day.has_activity:
                continue

            # Partial activity is usable evidence, but it cannot be reconstructed
            # safely without both day boundaries. It is handled by the session
            # as a non-blocking skipped day instead of invalidating the workbook.
            if day.has_partial_activity:
                continue

            if not day.events:
                issues.append(
                    f"{label}: día activo sin filas de detalle"
                )
                continue

            if day.events[-1].kind != "closing":
                issues.append(
                    f"{label}: la última fila de detalle no es closing"
                )

            if day.computed_driving_seconds != day.driving_seconds:
                issues.append(
                    f"{label}: conducción calculada "
                    f"{day.computed_driving_seconds}s != resumen "
                    f"{day.driving_seconds}s"
                )

            if day.computed_stop_seconds != day.stop_seconds:
                issues.append(
                    f"{label}: parada calculada "
                    f"{day.computed_stop_seconds}s != resumen "
                    f"{day.stop_seconds}s"
                )

            real_event_count = sum(
                1
                for event in day.events
                if event.kind != "opening"
            )

            distance_tolerance = self._distance_rounding_tolerance(
                real_event_count
            )

            if (
                abs(day.computed_distance_km - day.distance_km)
                > distance_tolerance
            ):
                issues.append(
                    f"{label}: distancia calculada "
                    f"{day.computed_distance_km:.3f} km != resumen "
                    f"{day.distance_km:.3f} km"
                )

        day_totals: dict[str, dict[str, float]] = defaultdict(
            lambda: {
                "jornada": 0.0,
                "driving": 0.0,
                "stop": 0.0,
                "distance": 0.0,
                "distance_parts": 0.0,
                "speed": 0.0,
            }
        )

        for day in days:
            aggregate = day_totals[day.vehicle]
            aggregate["jornada"] += day.jornada_seconds
            aggregate["driving"] += day.driving_seconds
            aggregate["stop"] += day.stop_seconds
            aggregate["distance"] += day.distance_km
            aggregate["distance_parts"] += 1.0
            aggregate["speed"] = max(
                aggregate["speed"],
                day.max_speed_kmh,
            )

        for total in totals:
            aggregate = day_totals.get(total.vehicle)

            if aggregate is None:
                issues.append(
                    f"{total.vehicle}: aparece en Totales pero no en SubTotales"
                )
                continue

            if int(aggregate["jornada"]) != total.jornada_seconds:
                issues.append(
                    f"{total.vehicle}: JORNADA total no cuadra"
                )

            if int(aggregate["driving"]) != total.driving_seconds:
                issues.append(
                    f"{total.vehicle}: H.CONDUCCIÓN total no cuadra"
                )

            if int(aggregate["stop"]) != total.stop_seconds:
                issues.append(
                    f"{total.vehicle}: H.STOP total no cuadra"
                )

            total_distance_tolerance = self._distance_rounding_tolerance(
                int(aggregate["distance_parts"])
            )

            if (
                abs(aggregate["distance"] - total.distance_km)
                > total_distance_tolerance
            ):
                issues.append(
                    f"{total.vehicle}: kilómetros totales no cuadran"
                )

            if abs(aggregate["speed"] - total.max_speed_kmh) > 0.001:
                issues.append(
                    f"{total.vehicle}: velocidad punta total no cuadra"
                )

        return issues

    def _block_date(self, text: str) -> date | None:
        match = _BLOCK_RE.match(text)

        if not match:
            return None

        return datetime.strptime(
            match.group(1),
            "%d-%m-%Y",
        ).date()

    def _optional_datetime(self, value: Any) -> datetime | None:
        text = self._text(value)

        if text in {"", "--", "??"}:
            return None

        return self._parse_datetime(value)

    def _parse_datetime(self, value: Any) -> datetime:
        if isinstance(value, datetime):
            return value

        text = self._text(value)

        for fmt in (
            "%d-%m-%Y %H:%M:%S",
            "%d/%m/%Y %H:%M:%S",
        ):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass

        raise ValueError(f"Fecha/hora no reconocida: {value!r}")

    def _parse_clock_time(self, value: Any) -> time:
        if isinstance(value, datetime):
            return value.time().replace(microsecond=0)

        if isinstance(value, time):
            return value.replace(microsecond=0)

        text = self._text(value)

        parts = text.split(":")
        if len(parts) != 3:
            raise ValueError(f"Hora no reconocida: {value!r}")

        hour, minute, second = (int(part) for part in parts)

        return time(
            hour=hour,
            minute=minute,
            second=second,
        )

    def _duration_seconds(self, value: Any) -> int | None:
        if value is None:
            return None

        if isinstance(value, timedelta):
            return int(round(value.total_seconds()))

        if isinstance(value, time):
            return (
                value.hour * 3600
                + value.minute * 60
                + value.second
            )

        if isinstance(value, (int, float)):
            if float(value) == 0.0:
                return 0

            if 0 < float(value) < 1:
                return int(round(float(value) * 86400))

        text = self._text(value)

        if text in {"", "--", "??"}:
            return None

        parts = text.split(":")
        if len(parts) != 3:
            raise ValueError(f"Duración no reconocida: {value!r}")

        hours, minutes, seconds = (int(part) for part in parts)

        if hours < 0 or minutes < 0 or seconds < 0:
            raise ValueError(f"Duración negativa no válida: {value!r}")

        if minutes >= 60 or seconds >= 60:
            raise ValueError(f"Duración no válida: {value!r}")

        return hours * 3600 + minutes * 60 + seconds

    def _number(self, value: Any) -> float | None:
        if value is None:
            return None

        if isinstance(value, bool):
            return float(int(value))

        if isinstance(value, (int, float)):
            return float(value)

        text = self._text(value)

        if text in {"", "--", "??"}:
            return None

        normalized = (
            text.replace(" ", "")
            .replace(".", "")
            .replace(",", ".")
        )

        try:
            return float(normalized)
        except ValueError as exc:
            raise ValueError(
                f"Número no reconocido: {value!r}"
            ) from exc

    def _split_stop_address(
        self,
        value: str,
    ) -> tuple[str | None, str | None]:
        if not value:
            return None, None

        match = re.match(r"^\(([PE])\)\s*(.*)$", value)

        if not match:
            return None, value

        stop_type = match.group(1)
        address = match.group(2).strip() or None

        return stop_type, address

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""

        return str(value).strip()