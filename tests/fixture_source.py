from __future__ import annotations

import os
import tempfile
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook


PRIVATE_FIXTURE_DIR = Path("data/private/fixtures")

SYNTHETIC_VEHICLE_A = "VEHICLE ALPHA"
SYNTHETIC_VEHICLE_B = "VEHICLE BETA"


def _set_headers(workbook: Workbook) -> None:
    totals = workbook["Totales"]
    subtotals = workbook["SubTotales"]
    detail = workbook["Detalle"]

    totals.append([])
    totals.append([])
    totals.append([])
    totals.append(
        [
            "VEHÍCULO",
            "JORNADA",
            "H.MOTOR",
            "H.STOP",
            "Vel.PUNTA",
            "Kms",
        ]
    )

    subtotals.append([])
    subtotals.append([])
    subtotals.append([])
    subtotals.append(
        [
            "VEHÍCULO",
            "INICIO",
            "FIN",
            "JORNADA",
            "H.MOTOR",
            "H.STOP",
            "Vel.PUNTA",
            "Kms",
        ]
    )

    detail.append([])
    detail.append([])
    detail.append([])
    detail.append(
        [
            "VEHÍCULO",
            "PARA",
            "ARRANCA",
            "JORNADA",
            "H.MOTOR",
            "H.STOP",
            "Vel.PUNTA",
            "Kms",
            "DIRECCIÓN PARADA",
            "MAPA",
        ]
    )


def _set_export_range(
    workbook: Workbook,
    start_text: str,
    end_text: str,
) -> None:
    value = (
        f"INICIO: {start_text} "
        f"FIN: {end_text}"
    )

    for name in ("Totales", "SubTotales", "Detalle"):
        workbook[name]["A2"] = value


def _map_link(
    sheet,
    row: int,
    target: str,
) -> None:
    cell = sheet.cell(row, 10)
    cell.value = "Mapa"
    cell.hyperlink = target


def _build_primary(path: Path) -> None:
    workbook = Workbook()

    totals = workbook.active
    totals.title = "Totales"

    subtotals = workbook.create_sheet(
        "SubTotales"
    )
    detail = workbook.create_sheet(
        "Detalle"
    )

    _set_headers(workbook)
    _set_export_range(
        workbook,
        "10-01-2026 00:00:00",
        "11-01-2026 23:59:00",
    )

    # ---------------------------------------------------------
    # TOTALES
    # ---------------------------------------------------------

    totals.append(
        [
            SYNTHETIC_VEHICLE_A,
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
        ]
    )

    totals.append(
        [
            SYNTHETIC_VEHICLE_B,
            "0:20:00",
            "0:15:00",
            "0:05:00",
            50,
            "10,0",
        ]
    )

    # ---------------------------------------------------------
    # SUBTOTALES
    # ---------------------------------------------------------

    subtotals.append(
        ["Inicio de bloque: 10-01-2026"]
    )

    subtotals.append(
        [
            SYNTHETIC_VEHICLE_A,
            "10-01-2026 08:00:00",
            "10-01-2026 09:30:00",
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
        ]
    )

    subtotals.append(
        [
            SYNTHETIC_VEHICLE_B,
            "10-01-2026 10:00:00",
            "10-01-2026 10:20:00",
            "0:20:00",
            "0:15:00",
            "0:05:00",
            50,
            "10,0",
        ]
    )

    subtotals.append(
        ["Inicio de bloque: 11-01-2026"]
    )

    subtotals.append(
        [
            SYNTHETIC_VEHICLE_A,
            "--",
            "--",
            "0:0:0",
            "0:0:0",
            "0:0:0",
            0,
            "0,0",
        ]
    )

    # ---------------------------------------------------------
    # DETALLE — VEHICLE ALPHA
    #
    # Incluye:
    # - opening con valores arrastrados;
    # - segmentos;
    # - closing;
    # - diferencia de 0,2 km por redondeo.
    # ---------------------------------------------------------

    detail.append(
        ["Inicio de bloque: 10-01-2026"]
    )

    detail.append(
        [
            SYNTHETIC_VEHICLE_A,
            "10-01-2026 08:00:00",
            "10-01-2026 09:30:00",
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
            None,
            None,
        ]
    )

    # Opening: 7,5 km / 55 km/h son arrastre y deben ignorarse.
    detail.append(
        [
            None,
            "??",
            "08:00:00",
            None,
            None,
            0,
            55,
            "7,5",
            None,
            "Mapa",
        ]
    )

    detail.append(
        [
            None,
            "08:10:00",
            "08:40:00",
            None,
            "00:10:00",
            "00:30:00",
            60,
            "5,0",
            "(P) Synthetic Street 1",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-1",
    )

    detail.append(
        [
            None,
            "08:55:00",
            "09:25:00",
            None,
            "00:15:00",
            "00:30:00",
            70,
            "10,0",
            "(E) Synthetic Street 2",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-2",
    )

    detail.append(
        [
            None,
            "09:30:00",
            "??",
            None,
            "00:05:00",
            "??",
            80,
            "5,0",
            "(P) Synthetic Street 3",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-3",
    )

    # ---------------------------------------------------------
    # DETALLE — VEHICLE BETA
    #
    # Simula apertura omitida con primera fila contaminada:
    # RAW = 14 km / 90 km/h
    # valor normalizado esperado = 4 km
    # ---------------------------------------------------------

    detail.append(
        [
            SYNTHETIC_VEHICLE_B,
            "10-01-2026 10:00:00",
            "10-01-2026 10:20:00",
            "0:20:00",
            "0:15:00",
            "0:05:00",
            50,
            "10,0",
            None,
            None,
        ]
    )

    detail.append(
        [
            None,
            "10:05:00",
            "10:10:00",
            None,
            "00:05:00",
            "00:05:00",
            90,
            "14,0",
            "(P) Synthetic Avenue 1",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/beta-1",
    )

    detail.append(
        [
            None,
            "10:20:00",
            "??",
            None,
            "00:10:00",
            "??",
            50,
            "6,0",
            "(P) Synthetic Avenue 2",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/beta-2",
    )

    # ---------------------------------------------------------
    # DÍA SIN ACTIVIDAD
    # ---------------------------------------------------------

    detail.append(
        ["Inicio de bloque: 11-01-2026"]
    )

    detail.append(
        [
            SYNTHETIC_VEHICLE_A,
            "--",
            "--",
            "0:0:0",
            "0:0:0",
            "0:0:0",
            0,
            "0,0",
            None,
            None,
        ]
    )

    workbook.save(path)
    workbook.close()


def _build_duplicate(path: Path) -> None:
    workbook = Workbook()

    totals = workbook.active
    totals.title = "Totales"

    subtotals = workbook.create_sheet(
        "SubTotales"
    )
    detail = workbook.create_sheet(
        "Detalle"
    )

    _set_headers(workbook)
    _set_export_range(
        workbook,
        "10-01-2026 00:00:00",
        "10-01-2026 23:59:00",
    )

    totals.append(
        [
            SYNTHETIC_VEHICLE_A,
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
        ]
    )

    subtotals.append(
        ["Inicio de bloque: 10-01-2026"]
    )

    subtotals.append(
        [
            SYNTHETIC_VEHICLE_A,
            "10-01-2026 08:00:00",
            "10-01-2026 09:30:00",
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
        ]
    )

    detail.append(
        ["Inicio de bloque: 10-01-2026"]
    )

    detail.append(
        [
            SYNTHETIC_VEHICLE_A,
            "10-01-2026 08:00:00",
            "10-01-2026 09:30:00",
            "1:30:00",
            "0:30:00",
            "1:00:00",
            80,
            "20,2",
            None,
            None,
        ]
    )

    detail.append(
        [
            None,
            "??",
            "08:00:00",
            None,
            None,
            0,
            55,
            "7,5",
            None,
            "Mapa",
        ]
    )

    detail.append(
        [
            None,
            "08:10:00",
            "08:40:00",
            None,
            "00:10:00",
            "00:30:00",
            60,
            "5,0",
            "(P) Synthetic Street 1",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-1",
    )

    detail.append(
        [
            None,
            "08:55:00",
            "09:25:00",
            None,
            "00:15:00",
            "00:30:00",
            70,
            "10,0",
            "(E) Synthetic Street 2",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-2",
    )

    detail.append(
        [
            None,
            "09:30:00",
            "??",
            None,
            "00:05:00",
            "??",
            80,
            "5,0",
            "(P) Synthetic Street 3",
            "Mapa",
        ]
    )
    _map_link(
        detail,
        detail.max_row,
        "https://example.test/map/alpha-3",
    )

    workbook.save(path)
    workbook.close()


@lru_cache(maxsize=1)
def synthetic_fixture_files() -> tuple[Path, ...]:
    root = (
        Path(tempfile.gettempdir())
        / "gps_analyst_v01_synthetic_fixtures"
    )

    root.mkdir(
        parents=True,
        exist_ok=True,
    )

    primary = root / "synthetic_primary.xlsx"
    duplicate = root / "synthetic_duplicate.xlsx"

    _build_primary(primary)
    _build_duplicate(duplicate)

    return (
        duplicate,
        primary,
    )


def fixture_files() -> tuple[Path, ...]:
    mode = os.environ.get(
        "GPS_ANALYST_TEST_DATA",
        "synthetic",
    ).strip().lower()

    if mode == "synthetic":
        return synthetic_fixture_files()

    if mode == "private":
        files = tuple(
            sorted(
                PRIVATE_FIXTURE_DIR.glob(
                    "*.xlsx"
                )
            )
        )

        if not files:
            raise AssertionError(
                "GPS_ANALYST_TEST_DATA=private "
                "pero no hay XLSX en "
                "data/private/fixtures."
            )

        return files

    raise ValueError(
        "GPS_ANALYST_TEST_DATA debe ser "
        "'synthetic' o 'private'."
    )